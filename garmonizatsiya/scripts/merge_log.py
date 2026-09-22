#!/usr/bin/env python3
"""Слияние выгрузки приложения (CSV) с листом «ЛогВремени».

Использование:
    python3 scripts/merge_log.py путь/к/timelog_clean.xlsx garmonizatsiya_ГГГГ-ММ-ДД.csv [выход.xlsx]

Исходный файл не трогается: результат пишется в выход.xlsx (по умолчанию
<имя>_merged.xlsx рядом с исходным). Проверьте его и замените исходный сами.

Запись ищется по паре «начало (до минуты) + акт»:
- есть и уже с концом — пропускается, поэтому повторная выгрузка безопасна;
- есть без конца, а в CSV конец указан — конец дописывается (так закрывается
  открытое действие из Excel и последнее «идущее» из прошлой выгрузки);
- нет — добавляется новой строкой: B акт, D сфера, G начало, H конец, J тип.
  Формулы остальных столбцов копируются из последней строки лога.
"""
import csv
import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

SHEET = "ЛогВремени"
COL_ACT, COL_SPHERE, COL_START, COL_END, COL_TYPE = 2, 4, 7, 8, 10   # B, D, G, H, J
OWN_COLS = {COL_ACT, COL_SPHERE, COL_START, COL_END, COL_TYPE}
FMT = "%d.%m.%Y %H:%M"


def read_csv(path):
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            if not (r.get("Акт") or "").strip() or not r.get("Начало"):
                continue
            rows.append({
                "act": r["Акт"].strip(),
                "sphere": (r.get("Сфера") or "").strip(),
                "start": dt.datetime.strptime(r["Начало"].strip(), FMT),
                "end": dt.datetime.strptime(r["Конец"].strip(), FMT) if (r.get("Конец") or "").strip() else None,
                "type": (r.get("Тип") or "основной").strip(),
            })
    return rows


def minute(d):
    return d.replace(second=0, microsecond=0)


def main():
    if len(sys.argv) not in (3, 4):
        sys.exit(__doc__)
    src, csv_path = Path(sys.argv[1]), Path(sys.argv[2])
    out = Path(sys.argv[3]) if len(sys.argv) == 4 else src.with_name(src.stem + "_merged.xlsx")
    if out.resolve() == src.resolve():
        sys.exit("Выходной файл совпадает с исходным — укажите другое имя.")

    incoming = read_csv(csv_path)
    if not incoming:
        sys.exit("В CSV нет записей.")

    wb = load_workbook(src)
    ws = wb[SHEET]

    index, last_row = {}, 1
    for row in range(2, ws.max_row + 1):
        start, act = ws.cell(row, COL_START).value, ws.cell(row, COL_ACT).value
        if isinstance(start, dt.datetime) and act:
            index[(minute(start), str(act).strip())] = row
            last_row = row

    added = closed = skipped = 0
    for r in incoming:
        row = index.get((minute(r["start"]), r["act"]))
        if row:
            if ws.cell(row, COL_END).value is None and r["end"]:
                ws.cell(row, COL_END).value = r["end"]
                closed += 1
            else:
                skipped += 1
            continue
        new = ws.max_row + 1
        for col in range(1, ws.max_column + 1):
            tpl = ws.cell(last_row, col)
            cell = ws.cell(new, col)
            if col not in OWN_COLS and isinstance(tpl.value, str) and tpl.value.startswith("="):
                cell.value = Translator(tpl.value, origin=tpl.coordinate).translate_formula(cell.coordinate)
            if tpl.has_style:
                cell._style = tpl._style
        ws.cell(new, COL_ACT).value = r["act"]
        ws.cell(new, COL_SPHERE).value = r["sphere"]
        ws.cell(new, COL_START).value = r["start"]
        ws.cell(new, COL_END).value = r["end"]
        ws.cell(new, COL_TYPE).value = r["type"]
        index[(minute(r["start"]), r["act"])] = new
        added += 1

    wb.save(out)
    print(f"Готово: {out} — добавлено {added}, закрыто {closed}, уже были {skipped}.")


if __name__ == "__main__":
    main()
