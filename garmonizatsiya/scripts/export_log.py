#!/usr/bin/env python3
"""Выгрузка лога времени из Excel в data/data.json для приложения.

Использование:
    python3 scripts/export_log.py путь/к/timelog_clean.xlsx

Читает лист «ЛогВремени»: B — акт, D — сфера, G — начало, H — конец, J — тип
(«основной» / «фон»). Время в Excel — московское, без часового пояса.
"""
import collections
import datetime as dt
import json
import sys
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

MSK = ZoneInfo("Europe/Moscow")
HISTORY_DAYS = 35      # сколько дней полных записей класть в приложение (окна нужд — до 7 дней, график — 14)
RATE_DAYS = 28         # окно для темпа навыков
MAX_INTERVAL_H = 14    # интервалы длиннее — остатки разрывов в логировании, не учитываются
TOP_ACTS = 120         # сколько актов отдавать в автоподсказку

OUT = Path(__file__).resolve().parent.parent / "data" / "data.json"


def load_rows(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["ЛогВремени"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        act, sphere, start, end, typ = r[1], r[3], r[6], r[7], r[9]
        if not isinstance(start, dt.datetime) or not act:
            continue
        rows.append({
            "act": str(act).strip(),
            "sphere": str(sphere or "").strip(),
            "start": start,
            "end": end if isinstance(end, dt.datetime) else None,
            "bg": (typ or "").strip() == "фон",
        })
    return rows


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    rows = load_rows(sys.argv[1])
    if not rows:
        sys.exit("В листе «ЛогВремени» не нашлось ни одной записи.")

    ts = lambda d: int(d.replace(tzinfo=MSK).timestamp())
    hrs = lambda r: (r["end"] - r["start"]).total_seconds() / 3600
    closed = [r for r in rows if r["end"]]
    main_rows = sorted((r for r in closed if not r["bg"] and 0 < hrs(r) <= MAX_INTERVAL_H),
                       key=lambda r: r["start"])
    opens = [r for r in rows if r["end"] is None and not r["bg"]]
    open_entry = max(opens, key=lambda r: r["start"]) if opens else None
    last = max(r["start"] for r in rows)

    sphere_names, act_names, sidx, aidx = [], [], {}, {}

    def si(s):
        if s not in sidx:
            sidx[s] = len(sphere_names); sphere_names.append(s)
        return sidx[s]

    def ai(a):
        if a not in aidx:
            aidx[a] = len(act_names); act_names.append(a)
        return aidx[a]

    since = last - dt.timedelta(days=HISTORY_DAYS)
    entries = [[ts(r["start"]), round(hrs(r), 3), si(r["sphere"]), ai(r["act"])]
               for r in main_rows if r["start"] >= since]

    tot_s, tot_a = collections.Counter(), collections.Counter()
    for r in main_rows:
        tot_s[r["sphere"]] += hrs(r); tot_a[r["act"]] += hrs(r)

    rate_since = last - dt.timedelta(days=RATE_DAYS)
    r_s, r_a = collections.Counter(), collections.Counter()
    for r in main_rows:
        if r["start"] >= rate_since:
            r_s[r["sphere"]] += hrs(r); r_a[r["act"]] += hrs(r)

    by_act = collections.defaultdict(collections.Counter)
    for r in closed:
        if not r["bg"]:
            by_act[r["act"]][r["sphere"]] += 1
    top = [a for a, _ in collections.Counter({a: sum(c.values()) for a, c in by_act.items()}).most_common(TOP_ACTS)]

    data = {
        "tz": "Europe/Moscow",
        "totalRows": len(rows),
        "sphereNames": sphere_names,
        "actNames": act_names,
        "entries": entries,
        "open": ({"t": ts(open_entry["start"]), "act": open_entry["act"], "sphere": open_entry["sphere"]}
                 if open_entry else None),
        "totalSphere": {k: round(v, 1) for k, v in tot_s.items()},
        "totalAct": {k: round(v, 1) for k, v in tot_a.most_common(80)},
        "rate28Sphere": {k: round(v, 2) for k, v in r_s.items()},
        "rate28Act": {k: round(v, 2) for k, v in r_a.most_common(80)},
        "actSphere": {a: by_act[a].most_common(1)[0][0] for a in top},
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Готово: {OUT} — {len(rows)} строк лога, {len(entries)} записей за {HISTORY_DAYS} дн., "
          f"последняя запись {last:%d.%m %H:%M}")


if __name__ == "__main__":
    main()
